from __future__ import annotations

import re
import sys
import time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from helper_for_6_imf_openai_helper import get_client
# from home_openai_helper import get_client


# ================================================================
# SETTINGS
# ================================================================

example_file: Path | None = Path(r"k:\example.txt")
stems_file = Path(r"k:\stems.txt")

model = "gpt-5.6-sol"

chart_dir = Path("k:/excel_and_png")
output_dir = Path("k:/r_code")

MAX_ATTEMPTS = 2
SAMPLE_ROWS = 12

client = get_client()

if not chart_dir.exists():
    raise FileNotFoundError(
        f"The chart directory was not found:\n{chart_dir}"
    )

output_dir.mkdir(parents=True, exist_ok=True)


# ================================================================
# TOKEN USAGE
# ================================================================

token_usage = {
    "input": 0,
    "output": 0,
    "total": 0,
}


# ================================================================
# PROMPT
# ================================================================

prompt = r"""
You are converting an Excel chart into executable R code.

You are given:

1. A PNG image showing the chart that must be reproduced.
2. An Excel chart-link workbook containing the underlying data and, where
   applicable, references to external source workbooks.
3. A VERIFIED WORKBOOK STRUCTURE extracted directly from the Excel file by
   Python.
4. Optionally, a text file containing example R code that shows the preferred
   coding and chart-formatting style.

Your task is to write concise R code that reproduces the chart as closely as
possible.

The VERIFIED WORKBOOK STRUCTURE is authoritative for:
- sheet names
- column counts
- column names
- observed column types
- sample values

The workbook is authoritative for the data.
The PNG is authoritative for chart appearance.


CRITICAL SOURCE RULES

* NEVER invent a workbook column, sheet name, or range.

* NEVER invent an index, x, quarter, date, category, label, series, Dummy1,
  Dummy2, or any other SOURCE variable unless that exact name appears in the
  VERIFIED WORKBOOK STRUCTURE.

* A source column may be referenced by name in generated R code ONLY if that
  exact name appears in the VERIFIED WORKBOOK STRUCTURE.

* Do not infer workbook column names from chart legend labels, titles, axis
  labels, captions, annotations, expected chart structure, or economic meaning.

* A plausible but unverified source-column name is an error.

* Preserve the actual workbook table structure.

* Never select more source columns than the verified workbook contains.

* Never assign more source-column names than the verified workbook contains.

* Derived R variables may be created AFTER reading verified source data, but
  they must be explicitly created in the R code.

* Do not treat a derived R variable as if it were a workbook column.


DATA TYPES AND ARITHMETIC

* The VERIFIED WORKBOOK STRUCTURE reports an observed type for every column.

* Before arithmetic, comparisons, transformations, scales, or statistics,
  check the verified observed type of every participating column.

* A column can contain values that LOOK numeric while being stored as text.
  If the verified structure says "text containing numeric values" or
  "mixed numeric and numeric-looking text", explicitly convert that column
  with as.numeric() before using it in arithmetic or numeric plotting.

* Never assume a column is numeric merely because its displayed values look
  like numbers.

* If several columns participate in arithmetic, ensure every participating
  column is numeric first.

* For example, if Dummy1 is numeric but Dummy2, Dummy3, and Dummy4 are reported
  as numeric-looking text, valid code is:

      plot_data <- read_excel(
        path = workbook_path,
        sheet = "data"
      ) %>%
        mutate(
          across(c(Dummy1, Dummy2, Dummy3, Dummy4), as.numeric),
          Andean_lower = Dummy1,
          Andean_upper = Dummy1 + Dummy2,
          CAPDR_lower = Dummy1 + Dummy2 + Dummy3,
          CAPDR_upper = Dummy1 + Dummy2 + Dummy3 + Dummy4
        )

* Do not convert descriptive text or category columns to numeric.


READING DATA

* Use readxl to read the workbook.

* Use verified sheet names exactly.

* Do not invent an Excel range.

* If no range is explicitly required, read the verified sheet directly.

* If an external workbook path, sheet, and range are explicitly given in
  workbook metadata, use those exact references.

* Otherwise use the supplied chart-link workbook.

* Use the full supplied Windows path.

* If source column names are blank, duplicated, or ambiguous, use column
  positions rather than inventing names.

* If readxl repairs blank or duplicated names into ...1, ...2, etc., either use
  those exact repaired names or select by position.

* Do not write runtime code that tries to discover or guess workbook structure.

* Do not use grep() searches for likely headers, fallback column names,
  candidate-name lists, heuristic schema discovery, or alternative guessed
  source structures.

The workbook structure has already been inspected by Python.


DATA TRANSFORMATION

* Use the simplest transformation consistent with the verified table.

* If the workbook is already a normal wide table with one x/category column
  and several value columns, pivot_longer() is normally sufficient.

* Do not create an artificial numeric x variable merely to make plotting easier.

* If the workbook x-axis column contains text such as 2022Q1, 2022Q2, 2022Q3,
  use that column directly.

* Preserve workbook ordering where necessary with factor levels.

* Do not split one source column into an invented x column plus another label
  column.

* Do not create x_breaks or x_labels unless they are genuinely needed.


X AXIS

Determine whether the actual x-axis variable is:
- discrete categorical
- numeric continuous
- Date
- POSIX date/time

Use:
- scale_x_discrete() for categorical text values
- scale_x_continuous() only for genuinely numeric x-values
- scale_x_date() only for Date values

Do not pass character breaks to scale_x_continuous().

If the x-axis consists of four-digit numeric years, format them as years.

If the workbook genuinely contains Excel serial dates, convert them correctly.


CHART

* Return only executable R code.
* Do not use Markdown fences.
* Do not include commentary or explanation.

* Use readxl and ggplot2.
* Use dplyr, tidyr, patchwork, or other common packages only when useful.
* Include all required library() calls.

* Reproduce the supplied PNG as closely as practical, including:
  - title
  - subtitle
  - legend text and order
  - colors
  - fills
  - line types and widths
  - markers
  - axis titles, limits, ticks, and labels
  - annotations
  - captions
  - panel arrangement
  - spacing
  - theme

* Match the visible x-axis range in the PNG where appropriate.

* Use a secondary axis only when clearly shown by the PNG or workbook metadata.

* If using a secondary axis, use sec_axis() with a mathematically consistent
  reversible transformation.

* Do not duplicate or contradict theme(), scale_*(), guides(), labs(), or
  coordinate settings.

* Prefer simple, direct R code.


EXAMPLE R CODE

If an example text file is attached:

* It contains R code even though its extension is .txt.
* Treat it as example R code.
* Use it as a guide for chart style, fonts, colors, spacing, legends, titles,
  captions, and ggplot conventions.
* Do not copy its data transformations unless they match the VERIFIED WORKBOOK
  STRUCTURE.
* The PNG remains authoritative for appearance.
* The VERIFIED WORKBOOK STRUCTURE remains authoritative for source structure
  and observed data types.


FINAL CHECK

Before returning R code:

1. Check every sheet name against the VERIFIED WORKBOOK STRUCTURE.
2. Check every source column name against the VERIFIED WORKBOOK STRUCTURE.
3. Confirm that no source variable has been invented.
4. Confirm that the number of selected source columns does not exceed the
   verified number.
5. Confirm that every derived variable is explicitly created after import.
6. Confirm that no artificial index/x source variable has been created when the
   workbook already contains the x/category column.
7. Confirm that categorical x-values have not been unnecessarily converted to
   numeric.
8. Check the observed type of every column used in arithmetic.
9. Convert numeric-looking text columns with as.numeric() before arithmetic.
10. Confirm that every plotted series corresponds to actual workbook data.
11. Confirm that generated R code is syntactically valid.
12. Confirm that theme(), scale_*(), and guide settings are not contradictory.


OUTPUT

* Assign the completed plot to an object named chart.
* Then execute print(chart).

* Create a "figures" subdirectory inside the supplied output directory if
  necessary.

* Save with ggsave() using:
      dpi = 600
      width = 5
      height = 5
      units = "in"

* Use the input chart stem as the output PNG filename.
* Use Windows-compatible paths.
* Preserve non-breaking spaces in supplied paths.

If the required workbook structure or chart-series mapping genuinely cannot be
determined, return exactly:

# FAILED
"""


# ================================================================
# HELPERS
# ================================================================

def remove_markdown_fences(text: str) -> str:
    result = text.strip()
    result = re.sub(r"^```[A-Za-z]*\s*\n?", "", result)
    result = re.sub(r"\n?```\s*$", "", result)
    return result.strip()


def response_is_bad(text: str) -> bool:
    stripped = text.strip()

    if not stripped or stripped == "# FAILED":
        return True

    head = stripped[:300].lower()

    refusal_phrases = (
        "cannot assist",
        "i'm sorry",
        "i am sorry",
        "i'm unable to",
        "i am unable to",
    )

    return any(phrase in head for phrase in refusal_phrases)


def validate_inputs(
    chart_stem: str,
    png_file: Path,
    excel_file: Path,
) -> None:
    missing_files = [
        path
        for path in (png_file, excel_file)
        if not path.exists()
    ]

    if missing_files:
        available_files = sorted(
            path.name
            for path in chart_dir.glob(f"{chart_stem}*")
        )

        available_text = (
            "\n".join(f"- {name}" for name in available_files)
            if available_files
            else "(No similarly named files were found.)"
        )

        missing_text = "\n".join(str(path) for path in missing_files)

        raise FileNotFoundError(
            "The following input files were not found:\n"
            f"{missing_text}\n\n"
            f"Resolved chart directory:\n{chart_dir}\n\n"
            f"Files beginning with '{chart_stem}':\n{available_text}"
        )

    if example_file is not None and not example_file.exists():
        raise FileNotFoundError(
            f"Example R file not found:\n{example_file}"
        )


def is_numeric_text(value: Any) -> bool:
    """
    True if value is a non-empty string that can be parsed as a number.
    """
    if not isinstance(value, str):
        return False

    text = value.strip()

    if not text:
        return False

    # Allow ordinary thousands separators.
    candidate = text.replace(",", "")

    try:
        float(candidate)
        return True
    except ValueError:
        return False


def classify_column(values: list[Any]) -> str:
    """
    Classify the observed nonblank Excel values in one column.
    """
    nonblank = [
        value
        for value in values
        if value is not None
        and not (
            isinstance(value, str)
            and value.strip() == ""
        )
    ]

    if not nonblank:
        return "blank"

    numeric_count = sum(
        isinstance(value, (int, float)) and not isinstance(value, bool)
        for value in nonblank
    )

    boolean_count = sum(
        isinstance(value, bool)
        for value in nonblank
    )

    text_values = [
        value
        for value in nonblank
        if isinstance(value, str)
    ]

    numeric_text_count = sum(
        is_numeric_text(value)
        for value in text_values
    )

    if numeric_count == len(nonblank):
        return "numeric"

    if boolean_count == len(nonblank):
        return "logical"

    if len(text_values) == len(nonblank):
        if numeric_text_count == len(nonblank):
            return "text containing numeric values"

        if numeric_text_count > 0:
            return "mixed text, including numeric-looking text"

        return "text"

    if numeric_count + numeric_text_count == len(nonblank):
        return "mixed numeric and numeric-looking text"

    return "mixed types"


def get_workbook_structure(
    excel_file: Path,
    sample_rows: int = SAMPLE_ROWS,
) -> str:
    """
    Inspect the workbook locally and return authoritative sheet names,
    dimensions, headers, observed column types, and sample values.
    """
    workbook = load_workbook(
        excel_file,
        read_only=True,
        data_only=True,
    )

    sections: list[str] = []

    try:
        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]

            rows = list(
                ws.iter_rows(values_only=True)
            )

            while rows and all(
                value is None
                for value in rows[-1]
            ):
                rows.pop()

            if not rows:
                sections.append(
                    "\n".join(
                        [
                            "=" * 60,
                            f"Sheet: {sheet_name}",
                            "Rows: 0",
                            "Columns: 0",
                            "(empty sheet)",
                        ]
                    )
                )
                continue

            max_col = 0

            for row in rows:
                for column_number, value in enumerate(row, start=1):
                    if value is not None:
                        max_col = max(max_col, column_number)

            if max_col == 0:
                continue

            rows = [
                tuple(row[:max_col])
                for row in rows
            ]

            headers = rows[0]

            section = [
                "=" * 60,
                f"Sheet: {sheet_name}",
                f"Rows including first row: {len(rows)}",
                f"Columns: {max_col}",
                "",
                "Columns exactly as stored, with observed types:",
            ]

            for column_number in range(1, max_col + 1):
                header = headers[column_number - 1]

                values = [
                    row[column_number - 1]
                    if column_number - 1 < len(row)
                    else None
                    for row in rows[1:]
                ]

                observed_type = classify_column(values)

                section.append(
                    f"{column_number}: "
                    f"name={repr(header)}; "
                    f"observed_type={observed_type}"
                )

            section.append("")
            section.append(
                f"First {min(sample_rows, len(rows))} rows exactly as stored:"
            )

            for row_number, row in enumerate(
                rows[:sample_rows],
                start=1,
            ):
                values = [
                    "" if value is None else str(value)
                    for value in row
                ]

                section.append(
                    f"Row {row_number}: "
                    + " | ".join(values)
                )

            sections.append(
                "\n".join(section)
            )

    finally:
        workbook.close()

    return "\n\n".join(sections)


def upload_file(path: Path) -> str:
    with path.open("rb") as file_handle:
        uploaded_file = client.files.create(
            file=file_handle,
            purpose="assistants",
        )

    return uploaded_file.id


def delete_uploaded_files(
    file_ids: list[str],
) -> None:
    for file_id in file_ids:
        try:
            client.files.delete(file_id)
        except Exception:
            pass


def update_token_usage(
    response: Any,
) -> None:
    usage = response.usage

    if usage is None:
        print("Token usage was not returned.")
        return

    input_tokens = usage.input_tokens or 0
    output_tokens = usage.output_tokens or 0
    total_tokens = usage.total_tokens or 0

    token_usage["input"] += input_tokens
    token_usage["output"] += output_tokens
    token_usage["total"] += total_tokens

    print()
    print("Token usage for this call:")
    print(f"  Input:  {input_tokens:,}")
    print(f"  Output: {output_tokens:,}")
    print(f"  Total:  {total_tokens:,}")

    print("Running token usage:")
    print(f"  Input:  {token_usage['input']:,}")
    print(f"  Output: {token_usage['output']:,}")
    print(f"  Total:  {token_usage['total']:,}")
    print()


def build_request_content(
    chart_stem: str,
    png_file: Path,
    excel_file: Path,
    uploaded_image_id: str,
    uploaded_excel_id: str,
    uploaded_example_id: str | None = None,
) -> list[dict[str, Any]]:
    workbook_structure = get_workbook_structure(
        excel_file
    )

    instruction_text = (
        f"{prompt}\n\n"

        "============================================================\n"
        "VERIFIED WORKBOOK STRUCTURE\n"
        "============================================================\n\n"

        f"{workbook_structure}\n\n"

        "============================================================\n"
        "IMPORTANT\n"
        "============================================================\n\n"

        "The workbook structure above was extracted directly from the Excel "
        "file by Python before this request was submitted. It is authoritative.\n\n"

        "Use ONLY source sheet names and source columns listed above.\n\n"

        "The observed_type field is authoritative evidence about the observed "
        "Excel values. If a column is reported as 'text containing numeric "
        "values' or 'mixed numeric and numeric-looking text', explicitly "
        "convert it to numeric in R before arithmetic or numeric plotting.\n\n"

        "Do NOT invent source columns such as index, x, quarter, date, "
        "category, series, label, Dummy1, Dummy2, lower, upper, or any other "
        "source variable unless that exact source name appears above.\n\n"

        "Derived variables may be created only AFTER verified source columns "
        "have been read, and the R code must explicitly create them.\n\n"

        "============================================================\n"
        "ATTACHED FILES\n"
        "============================================================\n\n"

        "ATTACHED CHART IMAGE:\n"
        f"- Filename: {png_file.name}\n\n"

        "ATTACHED EXCEL WORKBOOK:\n"
        f"- Filename: {excel_file.name}\n"
        f"- Full on-disk path: {excel_file}\n\n"

        "INPUT CHART STEM:\n"
        f"{chart_stem}\n\n"

        "OUTPUT R SCRIPT DIRECTORY:\n"
        f"{output_dir}"
    )

    if (
        uploaded_example_id is not None
        and example_file is not None
    ):
        instruction_text += (
            "\n\n"
            "ATTACHED EXAMPLE R FILE:\n"
            f"- Filename: {example_file.name}\n"
            "- This .txt file contains R code.\n"
            "- Treat its contents as R code.\n"
            "- Use it as an example of preferred chart formatting and coding style.\n"
            "- Do not copy its workbook structure or data transformations unless "
            "they match the VERIFIED WORKBOOK STRUCTURE."
        )

    content: list[dict[str, Any]] = [
        {
            "type": "input_text",
            "text": instruction_text,
        },
        {
            "type": "input_image",
            "file_id": uploaded_image_id,
            "detail": "high",
        },
        {
            "type": "input_file",
            "file_id": uploaded_excel_id,
        },
    ]

    if uploaded_example_id is not None:
        content.append(
            {
                "type": "input_file",
                "file_id": uploaded_example_id,
            }
        )

    return content


def save_error(
    error: Exception | str,
    error_file: Path,
) -> None:
    if isinstance(error, Exception):
        error_text = (
            f"{type(error).__name__}: "
            f"{error}"
        )
    else:
        error_text = error

    error_file.write_text(
        error_text,
        encoding="utf-8",
    )


# ================================================================
# MAIN
# ================================================================

def generate_chart(
    chart_stem: str,
) -> Path | None:
    png_file = chart_dir / f"{chart_stem}.png"
    excel_file = chart_dir / f"{chart_stem}.xlsx"

    output_file = output_dir / f"{chart_stem}.R"
    error_file = output_dir / f"{chart_stem}_error.txt"

    validate_inputs(
        chart_stem,
        png_file,
        excel_file,
    )

    print()
    print("=" * 70)
    print(f"Processing: {chart_stem}")
    print(f"PNG:        {png_file}")
    print(f"Workbook:   {excel_file}")
    print("=" * 70)

    workbook_structure = get_workbook_structure(
        excel_file
    )

    print()
    print("Verified workbook structure:")
    print(workbook_structure)
    print()

    uploaded_ids: list[str] = []

    try:
        uploaded_image_id = upload_file(
            png_file
        )

        uploaded_ids.append(
            uploaded_image_id
        )

        uploaded_excel_id = upload_file(
            excel_file
        )

        uploaded_ids.append(
            uploaded_excel_id
        )

        uploaded_example_id: str | None = None

        if example_file is not None:
            uploaded_example_id = upload_file(
                example_file
            )

            uploaded_ids.append(
                uploaded_example_id
            )

        content = build_request_content(
            chart_stem=chart_stem,
            png_file=png_file,
            excel_file=excel_file,
            uploaded_image_id=uploaded_image_id,
            uploaded_excel_id=uploaded_excel_id,
            uploaded_example_id=uploaded_example_id,
        )

        r_code = ""

        for attempt in range(
            1,
            MAX_ATTEMPTS + 1,
        ):
            print(
                f"Submitting request "
                f"(attempt {attempt}/{MAX_ATTEMPTS})..."
            )

            response = client.responses.create(
                model=model,
                max_output_tokens=16000,
                input=[
                    {
                        "role": "user",
                        "content": content,
                    }
                ],
            )

            update_token_usage(
                response
            )

            r_code = remove_markdown_fences(
                response.output_text or ""
            )

            if not response_is_bad(
                r_code
            ):
                break

            if attempt < MAX_ATTEMPTS:
                time.sleep(2)

        if response_is_bad(
            r_code
        ):
            save_error(
                r_code or "(empty response)",
                error_file,
            )

            print(
                f"Generation failed: {chart_stem}"
            )

            return None

        output_file.write_text(
            r_code.rstrip() + "\n",
            encoding="utf-8",
        )

        print(
            f"Saved R code to: {output_file}"
        )

        return output_file

    except Exception as error:
        save_error(
            error,
            error_file,
        )

        print(
            f"ERROR: {chart_stem}: "
            f"{type(error).__name__}: {error}"
        )

        raise

    finally:
        delete_uploaded_files(
            uploaded_ids
        )


def read_stems_file() -> list[str]:
    """
    Read one chart stem per line from K:\\stems.txt.
    Blank lines and lines beginning with # are ignored.
    """
    if not stems_file.exists():
        raise FileNotFoundError(
            f"Stems file not found:\n{stems_file}"
        )

    stems: list[str] = []

    for line in stems_file.read_text(
        encoding="utf-8"
    ).splitlines():
        stem = line.strip()

        if not stem:
            continue

        if stem.startswith("#"):
            continue

        stems.append(
            stem
        )

    if not stems:
        raise ValueError(
            f"No stems found in:\n{stems_file}"
        )

    return stems


def resolve_stems(
    args: list[str],
) -> list[str]:
    """
    With no command-line arguments:
        read K:\\stems.txt

    With explicit stems:
        use those stems

    With --all:
        process every matching PNG/XLSX pair
    """
    if args == ["--all"]:
        stems = sorted(
            png.stem
            for png in chart_dir.glob("*.png")
            if (
                chart_dir
                / f"{png.stem}.xlsx"
            ).exists()
        )

        if not stems:
            raise FileNotFoundError(
                f"No PNG/xlsx pairs found in {chart_dir}"
            )

        return stems

    if args:
        return args

    return read_stems_file()


def print_final_token_usage() -> None:
    print()
    print("=" * 70)
    print("TOTAL TOKEN USAGE")
    print("=" * 70)

    print(
        f"Input tokens:  "
        f"{token_usage['input']:,}"
    )

    print(
        f"Output tokens: "
        f"{token_usage['output']:,}"
    )

    print(
        f"Total tokens:  "
        f"{token_usage['total']:,}"
    )

    print("=" * 70)


if __name__ == "__main__":
    try:
        stems = resolve_stems(
            sys.argv[1:]
        )

        print(
            f"Loaded {len(stems)} stems:"
        )

        for stem in stems:
            print(
                f"  {stem}"
            )

        for stem in stems:
            try:
                print()
                print(
                    f"STARTING: {stem}"
                )

                generate_chart(
                    stem
                )

            except Exception as error:
                print(
                    f"ERROR processing {stem}: "
                    f"{type(error).__name__}: "
                    f"{error}"
                )

    finally:
        print_final_token_usage()
