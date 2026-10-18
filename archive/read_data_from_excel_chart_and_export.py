import zipfile
import re
import os
import xml.etree.ElementTree as ET
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries, get_column_letter



NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a":   "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c":   "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "r":   "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "s":   "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
}


def _rels(z, part):
    """Read the .rels file for a part, return {rId: target path}."""
    folder, name = part.rsplit("/", 1)
    rels_path = f"{folder}/_rels/{name}.rels"
    out = {}
    for rel in ET.fromstring(z.read(rels_path)).findall("rel:Relationship", NS):
        target = rel.get("Target")
        if target.startswith("../"):
            target = "xl/" + target[3:]
        elif not target.startswith("/"):
            target = f"{folder}/{target}"
        out[rel.get("Id")] = target.lstrip("/")
    return out


def find_chart_part(z, sheet_name, chart_name):
    """Return the path of the chart XML (e.g. xl/charts/chart52.xml)."""
    wbxml = ET.fromstring(z.read("xl/workbook.xml"))
    wb_rels = _rels(z, "xl/workbook.xml")
    sheet_part = None
    for sh in wbxml.find("s:sheets", NS):
        if sh.get("name") == sheet_name:
            sheet_part = wb_rels[sh.get(f"{{{NS['r']}}}id")]
    if sheet_part is None:
        raise ValueError(f"No sheet named '{sheet_name}'")

    sheet_rels = _rels(z, sheet_part)
    drawing_part = next(t for t in sheet_rels.values()
                        if "drawings/drawing" in t and t.endswith(".xml"))
    drawing_rels = _rels(z, drawing_part)

    root = ET.fromstring(z.read(drawing_part))
    names = []
    for frame in root.iter(f"{{{NS['xdr']}}}graphicFrame"):
        cnvpr = frame.find(".//xdr:cNvPr", NS)
        names.append(cnvpr.get("name"))
        if cnvpr.get("name") == chart_name:
            chart_el = frame.find(".//c:chart", NS)
            return drawing_rels[chart_el.get(f"{{{NS['r']}}}id")]
    raise ValueError(f"No chart named '{chart_name}'. Found: {names}")


def _read_one_chart(filename, wb, sheet_name, chart_name, verbose=False):
    """Extract one chart's data into a DataFrame.

    Supports:
    - category charts using c:cat and c:val
    - XY scatter charts using c:xVal and c:yVal

    Cells in hidden rows or columns are excluded.
    """
    with zipfile.ZipFile(filename) as z:
        chart_part = find_chart_part(z, sheet_name, chart_name)
        chart_xml = ET.fromstring(z.read(chart_part))

    def _visible(ws, r, c):
        rd = ws.row_dimensions.get(r)
        if rd is not None and rd.hidden:
            return False

        cd = ws.column_dimensions.get(get_column_letter(c))
        if cd is not None and cd.hidden:
            return False

        return True

    def read_ref(f, visible_only=True):
        """Resolve an Excel chart formula to worksheet values.

        Handles formulas such as:
            Data_median!$A$2:$A$88
            'Sheet with spaces'!$B$2:$B$13
            'John''s Data'!$B$2:$B$13
        """
        if not f:
            return None

        match = re.match(
            r"(?:'((?:[^']|'')+)'|([^!]+))!(.+)",
            f
        )

        if match is None:
            raise ValueError(f"Cannot parse chart reference: {f}")

        sheet = match.group(1) or match.group(2)
        sheet = sheet.replace("''", "'")
        cell_range = match.group(3).replace("$", "")

        ws = wb[sheet]

        min_c, min_r, max_c, max_r = range_boundaries(cell_range)

        return [
            ws.cell(row=r, column=c).value
            for r in range(min_r, max_r + 1)
            for c in range(min_c, max_c + 1)
            if not visible_only or _visible(ws, r, c)
        ]

    def ref_of(el):
        """Return the formula inside numRef or strRef."""
        if el is None:
            return None

        for tag in ("c:numRef", "c:strRef"):
            ref = el.find(tag, NS)

            if ref is not None:
                formula = ref.find("c:f", NS)
                if formula is not None:
                    return formula.text

        return None

    def cache_of(el):
        """Return cached chart values."""
        if el is None:
            return None

        cache_tags = (
            "c:numRef/c:numCache",
            "c:strRef/c:strCache",
            "c:numLit",
            "c:strLit",
        )

        for tag in cache_tags:
            cache = el.find(tag, NS)

            if cache is None:
                continue

            points = {}

            for point in cache.findall("c:pt", NS):
                value = point.find("c:v", NS)

                points[int(point.get("idx"))] = (
                    value.text if value is not None else None
                )

            point_count = cache.find("c:ptCount", NS)

            if point_count is not None:
                n = int(point_count.get("val"))
            else:
                n = max(points, default=-1) + 1

            return [points.get(i) for i in range(n)]

        return None

    def values_of(el, visible_only=True):
        """Read referenced cells first, then use cached values."""
        if el is None:
            return None

        formula = ref_of(el)

        if formula:
            return read_ref(formula, visible_only=visible_only)

        return cache_of(el)

    def series_name(ser, i):
        """Extract the series name."""
        tx = ser.find("c:tx", NS)

        if tx is None:
            return f"series_{i}"

        formula = ref_of(tx)

        if formula:
            values = read_ref(formula, visible_only=False)

            if values and values[0] is not None:
                return str(values[0])

        literal = tx.find(".//c:v", NS)

        if literal is not None:
            return literal.text

        return f"series_{i}"

    def unique_name(name, existing_names):
        """Ensure that every output column name is unique."""
        original = name
        counter = 1

        while name in existing_names:
            counter += 1
            name = f"{original} ({counter})"

        return name

    series = list(chart_xml.iter(f"{{{NS['c']}}}ser"))

    if verbose:
        print(f"{chart_name}: {len(series)} series in chart XML")

    data = {}
    categories = None
    is_scatter = False

    for i, ser in enumerate(series):
        name = series_name(ser, i)

        # XY scatter charts use xVal and yVal.
        x_element = ser.find("c:xVal", NS)
        y_element = ser.find("c:yVal", NS)

        if x_element is not None or y_element is not None:
            is_scatter = True

            x_values = values_of(x_element) or []
            y_values = values_of(y_element) or []

            n_series = max(len(x_values), len(y_values))

            x_values += [None] * (n_series - len(x_values))
            y_values += [None] * (n_series - len(y_values))

            x_name = unique_name(f"{name}_x", data)
            data[x_name] = x_values

            y_name = unique_name(f"{name}_y", data)
            data[y_name] = y_values

        else:
            # Line, column, area and similar category charts.
            if categories is None:
                category_element = ser.find("c:cat", NS)
                categories = values_of(category_element)

            value_element = ser.find("c:val", NS)
            values = values_of(value_element)

            if values is not None:
                name = unique_name(name, data)
                data[name] = values

    # Pad all columns to the same length.
    n = max((len(values) for values in data.values()), default=0)

    data = {
        name: values + [None] * (n - len(values))
        for name, values in data.items()
    }

    df = pd.DataFrame(data)

    # Category charts generally have one common category column.
    # Scatter charts can have a different x range for every series.
    if not is_scatter and categories is not None:
        categories = categories + [None] * (n - len(categories))
        df.insert(0, "category", categories[:n])

    return df

def read_chart_data(filename, sheet_name, chart_name, chart_name_right=None,
                    output_name=None, out_dir=r"c:\test", verbose=False):
    """Read one chart (saved to sheet 'data') or a left/right pair
    (sheets 'datal' and 'datar') and write to an Excel file in out_dir.
    Hidden rows and columns are excluded from the exported data.

    Returns one DataFrame, or a (left, right) tuple when chart_name_right
    is given. Writes only when output_name is provided."""
    wb = load_workbook(filename, data_only=True)

    dfl = _read_one_chart(filename, wb, sheet_name, chart_name, verbose)
    dfr = (None if chart_name_right is None
           else _read_one_chart(filename, wb, sheet_name, chart_name_right,
                                verbose))

    if output_name:
        os.makedirs(out_dir, exist_ok=True)
        if not output_name.lower().endswith(".xlsx"):
            output_name += ".xlsx"
        out_path = os.path.join(out_dir, output_name)
        with pd.ExcelWriter(out_path) as xw:
            if dfr is None:
                dfl.to_excel(xw, sheet_name="data", index=False)
            else:
                dfl.to_excel(xw, sheet_name="datal", index=False)
                dfr.to_excel(xw, sheet_name="datar", index=False)
        if verbose:
            print(f"saved -> {out_path}")

    return dfl if dfr is None else (dfl, dfr)



if __name__ == "__main__":
    # dfl, dfr = read_chart_data(r"l:\Global Growth and Inflation.xlsx", "Chart",
    #                            "Chart 6", chart_name_right="Chart 22",
    #                            output_name="global growth and inflation",
    #                            out_dir=r"c:\test", verbose=True)
    # dfl, dfr = read_chart_data(r"\\DATA2\WHD\DATA\RS\REO\2026 Fall\Charts\Financial Markets Monitor.xlsm", "Chart",
    #                            "Chart 44", chart_name_right="Chart 45",
    #                            output_name="fig-2-3.xlsx",
    #                            out_dir=r"c:\test", verbose=True)



    # dfl = read_chart_data(r"\\DATA2\WHD\DATA\RS\REO\2026 Fall\Charts\Structural Primary Balance.xlsx", "Chart",
    #                            "Chart 4",    output_name="fig-9-2.xlsx",
    #                            out_dir=r"c:\test", verbose=True)

    dfl = read_chart_data(r"\\DATA2\WHD\DATA\RS\REO\2026 Fall\Charts\Pass-through_KDE_chart.xlsx", 
                            sheet_name="Chart",
                            chart_name="Chart 7",
                            output_name="fig-4-2.xlsx",
                            out_dir=r"K:\excel_and_png", verbose=True)

    