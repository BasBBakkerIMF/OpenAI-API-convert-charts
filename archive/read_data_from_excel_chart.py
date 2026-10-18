import zipfile
import re
import xml.etree.ElementTree as ET
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a":   "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c":   "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "r":   "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "s":   "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
}


def _rels(z, part):
    """Read the .rels file for a part, return {rId: target path}."""
    folder, name = part.rsplit("/", 1)
    rels_path = f"{folder}/_rels/{name}.rels"
    out = {}
    for rel in ET.fromstring(z.read(rels_path)).findall("rel:Relationship", NS):
        target = rel.get("Target")
        if target.startswith("../"):
            target = "xl/" + target[3:]
        elif not target.startswith("/"):
            target = f"{folder}/{target}"
        out[rel.get("Id")] = target.lstrip("/")
    return out


def find_chart_part(z, sheet_name, chart_name):
    """Return the path of the chart XML (e.g. xl/charts/chart52.xml)."""
    # sheet name -> sheet part
    wbxml = ET.fromstring(z.read("xl/workbook.xml"))
    wb_rels = _rels(z, "xl/workbook.xml")
    sheet_part = None
    for sh in wbxml.find("s:sheets", NS):
        if sh.get("name") == sheet_name:
            sheet_part = wb_rels[sh.get(f"{{{NS['r']}}}id")]
    if sheet_part is None:
        raise ValueError(f"No sheet named '{sheet_name}'")

    # sheet -> drawing part
    sheet_rels = _rels(z, sheet_part)
    drawing_part = next(t for t in sheet_rels.values() if "drawing" in t)
    drawing_rels = _rels(z, drawing_part)

    # drawing: find the graphicFrame with the right name, get its chart rId
    root = ET.fromstring(z.read(drawing_part))
    names = []
    for frame in root.iter(f"{{{NS['xdr']}}}graphicFrame"):
        cnvpr = frame.find(".//xdr:cNvPr", NS)
        names.append(cnvpr.get("name"))
        if cnvpr.get("name") == chart_name:
            chart_el = frame.find(".//c:chart", NS)
            return drawing_rels[chart_el.get(f"{{{NS['r']}}}id")]
    raise ValueError(f"No chart named '{chart_name}'. Found: {names}")


def read_chart_data(filename, sheet_name, chart_name, verbose=False):
    """Read the source data of a chart into a DataFrame, by internal
    Excel chart name (e.g. 'Chart 52'). Falls back to the values cached
    in the chart XML when a series has no live cell reference."""
    with zipfile.ZipFile(filename) as z:
        chart_part = find_chart_part(z, sheet_name, chart_name)
        chart_xml = ET.fromstring(z.read(chart_part))

    wb = load_workbook(filename, data_only=True)

    def read_ref(f):
        """Resolve 'Sheet1'!$B$2:$B$13 to a list of cell values."""
        m = re.match(r"'?([^'!]+)'?!(.+)", f)
        ws = wb[m.group(1)]
        min_c, min_r, max_c, max_r = range_boundaries(m.group(2).replace("$", ""))
        return [ws.cell(row=r, column=c).value
                for r in range(min_r, max_r + 1)
                for c in range(min_c, max_c + 1)]

    def ref_of(el):
        """Formula inside a numRef or strRef child, if any."""
        if el is None:
            return None
        for tag in ("c:numRef", "c:strRef"):
            ref = el.find(tag, NS)
            if ref is not None:
                return ref.find("c:f", NS).text
        return None

    def cache_of(el):
        """Plotted values from numCache/strCache when no cell ref exists."""
        if el is None:
            return None
        for tag in ("c:numRef/c:numCache", "c:strRef/c:strCache",
                    "c:numLit", "c:strLit"):
            cache = el.find(tag, NS)
            if cache is not None:
                pts = {int(pt.get("idx")): pt.find("c:v", NS).text
                       for pt in cache.findall("c:pt", NS)}
                n = int(cache.find("c:ptCount", NS).get("val"))
                return [pts.get(i) for i in range(n)]
        return None

    sers = list(chart_xml.iter(f"{{{NS['c']}}}ser"))
    if verbose:
        print(f"{len(sers)} series in chart XML")

    data, categories = {}, None
    for i, ser in enumerate(sers):
        # series name
        tx = ser.find("c:tx", NS)
        name = f"series_{i}"
        if tx is not None:
            f = ref_of(tx)
            if f:
                name = read_ref(f)[0]
            elif tx.find(".//c:v", NS) is not None:
                name = tx.find(".//c:v", NS).text

        # dedupe: never let two series share a key
        base, k = name, 1
        while name in data:
            k += 1
            name = f"{base} ({k})"

        # categories (from the first series that has them)
        if categories is None:
            cat = ser.find("c:cat", NS)
            f = ref_of(cat)
            categories = read_ref(f) if f else cache_of(cat)

        # values: live cell reference first, cached values as fallback
        val = ser.find("c:val", NS)
        f = ref_of(val)
        vals = read_ref(f) if f else cache_of(val)
        if vals is not None:
            data[name] = vals

    # pad to equal length so mismatched series don't break the constructor
    n = max((len(v) for v in data.values()), default=0)
    data = {k: v + [None] * (n - len(v)) for k, v in data.items()}

    df = pd.DataFrame(data)
    if categories is not None:
        categories = categories + [None] * (n - len(categories))
        df.insert(0, "category", categories[:n] if n else categories)
    return df


if __name__ == "__main__":
    # df = read_chart_data(r"l:\output gaps.xlsx", "Chart", "Chart 70",
    #                      verbose=True)

    # df = read_chart_data(r"l:\commodity prices.xlsx", "Chart", "Chart 12",
    #                          verbose=True)

    dfl= read_chart_data(r"l:\Global Growth and Inflation.xlsx", "Chart", "Chart 6",
                             verbose=True)
    dfr= read_chart_data(r"l:\Global Growth and Inflation.xlsx", "Chart", "Chart 22",
                             verbose=True)
    #print(df)